"""Защита API, проверка загрузок, настройки провайдеров, Excel в песочнице."""
import io

import pytest
from fastapi.testclient import TestClient

from app import config
from app.main import app, limiter
from app.sandbox import Sandbox

client = TestClient(app)


def test_token_required_when_configured(monkeypatch):
    monkeypatch.setattr(config, "APP_ACCESS_TOKEN", "s3cret")
    assert client.get("/api/state?events=0").status_code == 401
    assert client.get("/api/state?events=0", headers={"X-Frank-Token": "wrong"}).status_code == 401
    assert client.get("/api/state?events=0", headers={"X-Frank-Token": "s3cret"}).status_code == 200
    assert client.get("/api/health").json() == {"ok": True, "auth_required": True}


def test_rate_limit_on_expensive_endpoints():
    limiter._hits.clear()
    codes = [client.post("/api/recommend", json={"text": "x"}).status_code for _ in range(62)]
    assert codes[:60] == [200] * 60 and codes[-1] == 429
    limiter._hits.clear()


@pytest.mark.parametrize("name,data,code", [
    ("evil.exe", b"MZ....", 415),
    ("fake.pdf", b"not a pdf", 415),
    ("sheet.xlsx", b"not a zip", 415),
    ("big.txt", "BIG", 413),
    ("ok.pdf", b"%PDF-1.4 hello", 200),
    ("ok.csv", b"a,b\n1,2\n", 200),
], ids=["exe", "fake-pdf", "fake-xlsx", "too-big", "pdf", "csv"])
def test_upload_validation(name, data, code):
    limiter._hits.clear()
    if data == "BIG":
        data = b"x" * (config.MAX_UPLOAD_MB * 1024 * 1024 + 1)
    assert client.post("/api/upload", files={"file": (name, data)}).status_code == code


def test_settings_are_validated_and_persisted():
    r = client.get("/api/settings").json()
    assert "claude-sonnet-4-5" in r["options"]["models"] and "elevenlabs" in r["options"]["llm"]
    assert all(isinstance(v, bool) for v in r["keys"].values())            # наружу — только факт наличия ключа, не сам ключ
    assert client.put("/api/settings", json={"llm_model": "gpt-4.1"}).json()["current"]["llm_model"] == "gpt-4.1"
    assert client.get("/api/settings").json()["current"]["llm_model"] == "gpt-4.1"
    assert client.put("/api/settings", json={"llm_model": "my-private-model"}).status_code == 400
    assert client.put("/api/settings", json={"api_key": "x"}).status_code == 400


def test_balance_uses_tracked_usage_only():
    b = client.get("/api/balance").json()
    assert b["elevenlabs"] is None                 # без ключа — никаких выдуманных балансов
    assert set(b["usage"]) >= {"total", "by_model", "by_monster", "by_task", "recent", "by_kind"}


def test_transcribe_without_key_is_clear():
    limiter._hits.clear()
    r = client.post("/api/transcribe", files={"file": ("a.webm", b"x" * 500, "audio/webm")}, data={"lang": "en"})
    assert r.status_code in (409, 503)


def test_openpyxl_works_inside_the_sandbox():
    tests = """
import unittest, os, tempfile
class T(unittest.TestCase):
    def test_xlsx_roundtrip(self):
        from openpyxl import Workbook, load_workbook
        d = tempfile.mkdtemp()
        wb = Workbook(); ws = wb.active; ws.append(['amount']); ws.append([125.5]); wb.save(os.path.join(d, 't.xlsx'))
        self.assertEqual(load_workbook(os.path.join(d, 't.xlsx')).active['A2'].value, 125.5)
"""
    r = Sandbox().run_tests(code="def run(inp):\n    return {}\n", tests=tests, dependencies={})
    assert r.report["passed"] == 1, (r.report, r.stderr[-400:])
